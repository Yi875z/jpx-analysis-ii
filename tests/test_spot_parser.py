"""JPX現物（投資部門別売買状況）の新形式対応の回帰テスト。

2026-09-29 掲載分から JPX がファイル名（stock_val_1_YYMMNN.xls → stock_1_w_YYYYMMDD_YYYYMMDD.xlsx）と
Excel の配置（市場ごとのシート → 1シート・横持ち）を変え、9/18週・9/25週のレポートが欠落した。
守りたいこと:
  - 新形式の内訳を合算した値が、旧形式の「自己計・個人・海外投資家」と同じ定義になる
    （JPX サンプル＝2026年4月第1週を、旧形式から取り込んだ DB 値と突合）
  - 見出しや単位が想定と違えば、推測で読まずに例外で止まる
  - ファイル名から対象週を読む（新形式）。旧形式はページの期間表記で読む
  - JPX に掲載中なのにレポートが無い前の週を検出する

フィクスチャ:
  stock_1_w_20260924_20260925.xlsx   … 実ファイル（9/24〜9/25週、千円単位）
  stock_1_w_sample_2026041_yen.xlsx  … JPX のフォーマット変更告知のサンプル（2026年4月第1週）。
                                        見出しは「千円」だが中身は円単位（実ファイルの1000倍）。
"""

from datetime import date
from pathlib import Path

import openpyxl
import pytest

from scripts.jpx_week_resolver import (
    spot_files_by_week,
    spot_week_end_from_filename,
    unreported_weeks,
)
from scripts.parse_spot_xls import SPOT_INVESTORS, parse_spot_xls, parse_spot_xlsx_v3

FIXTURES = Path(__file__).parent / "fixtures"
REAL_0925 = FIXTURES / "stock_1_w_20260924_20260925.xlsx"
SAMPLE_YEN = FIXTURES / "stock_1_w_sample_2026041_yen.xlsx"

# 旧形式（stock_val_1_260401 相当）から取り込み済みの DB weekly_spot 2026-04-03 週（億円）。
# (buy, sell, net)
OLD_FORMAT_20260403 = {
    "dealer":     (48882.95, 64867.24, -15984.28),
    "individual": (97069.48, 101465.58, -4396.10),
    "foreign":    (265714.91, 246565.01, 19149.89),
    "inv_trust":  (12690.33, 11659.24, 1031.09),
    "corporate":  (3251.73, 3022.71, 229.02),
    "trust_bank": (9139.92, 7713.82, 1426.10),
}


def _by_type(rows):
    return {r["investor_type"]: r for r in rows}


def _copy_scaled(src: Path, dst: Path, divisor: int) -> Path:
    """数値セル（8行目以降）を divisor で割ったコピーを作る（円 → 千円 の換算用）。"""
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=8):
        for c in row:
            if isinstance(c.value, int):
                c.value = c.value // divisor
    wb.save(dst)
    return dst


def test_sample_sums_match_old_format_definition(tmp_path):
    """内訳の合算（自己=現金+信用 / 個人=現金+信用 / 海外=法人+個人）が旧形式の値と一致する。"""
    scaled = _copy_scaled(SAMPLE_YEN, tmp_path / "sample_thousand_yen.xlsx", 1000)
    got = _by_type(parse_spot_xls(str(scaled), "2026-04-03"))
    assert set(got) == set(SPOT_INVESTORS)
    for inv, (buy, sell, net) in OLD_FORMAT_20260403.items():
        assert got[inv]["buy_amount"] == pytest.approx(buy, abs=0.01), inv
        assert got[inv]["sell_amount"] == pytest.approx(sell, abs=0.01), inv
        assert got[inv]["net_amount"] == pytest.approx(net, abs=0.01), inv


def test_real_file_values():
    """実ファイル（9/24〜9/25週）の二市場・金額を、セルの値から手計算した値と照合する。"""
    got = _by_type(parse_spot_xls(str(REAL_0925), "2026-09-25"))
    assert set(got) == set(SPOT_INVESTORS)
    # 海外投資家 = 法人(T15/U15) + 個人(X15/Y15)
    #   売 12,451,103,435 + 37,916,041 = 12,489,019,476 千円
    #   買 12,358,355,242 + 37,656,433 = 12,396,011,675 千円
    assert got["foreign"]["sell_amount"] == pytest.approx(124890.19, abs=0.01)
    assert got["foreign"]["buy_amount"] == pytest.approx(123960.12, abs=0.01)
    assert got["foreign"]["net_amount"] == pytest.approx(-930.08, abs=0.01)
    # 自己 = 現金(D15/E15) + 信用(H15/I15)
    #   売 1,626,287,337 + 53,192,112 / 買 2,331,001,558 + 522,666
    assert got["dealer"]["net_amount"] == pytest.approx(6520.45, abs=0.01)
    # 信託銀行 = AZ15/BA15 のみ
    assert got["trust_bank"]["sell_amount"] == pytest.approx(6227.88, abs=0.01)
    assert got["trust_bank"]["net_amount"] == pytest.approx(-3673.01, abs=0.01)
    assert all(r["week_date"] == "2026-09-25" for r in got.values())


def test_unit_mismatch_is_rejected():
    """円単位のサンプルをそのまま読むと1000倍になる → 単位の妥当性チェックで止まる。"""
    with pytest.raises(ValueError, match="単位"):
        parse_spot_xlsx_v3(str(SAMPLE_YEN), "2026-04-03")


def test_header_change_is_rejected(tmp_path):
    """見出しが想定と違えば（JPX がまた配置を変えたら）推測で読まずに止まる。"""
    wb = openpyxl.load_workbook(REAL_0925)
    wb.worksheets[0]["AZ6"] = "その他金融機関 Other Financials"
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    with pytest.raises(ValueError, match="見出し"):
        parse_spot_xls(str(broken), "2026-09-25")


def test_column_shift_is_caught_by_balance_check(tmp_path):
    """差引列が 買−売 と合わない（列ずれ・値の取り違え）なら止まる。"""
    wb = openpyxl.load_workbook(REAL_0925)
    wb.worksheets[0]["U15"] = wb.worksheets[0]["U15"].value + 1
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    with pytest.raises(ValueError, match="差引"):
        parse_spot_xls(str(broken), "2026-09-25")


@pytest.mark.parametrize("url, expected", [
    ("https://www.jpx.co.jp/x/t13vrt0000026acb-att/stock_1_w_20260924_20260925.xlsx", date(2026, 9, 25)),
    ("https://www.jpx.co.jp/x/stock_1_w_20260914_20260918.xlsx", date(2026, 9, 18)),
    ("https://www.jpx.co.jp/x/stock_val_1_260902.xls", None),            # 旧形式はページ側で解決
    ("https://www.jpx.co.jp/x/stock_1_w_YYYYMMDD_YYYYMMDD.xlsx", None),  # 告知の見本
])
def test_spot_week_end_from_filename(url, expected):
    assert spot_week_end_from_filename(url) == expected


# 2026-10-02 時点の JPX 現物ページの構造を縮めたもの（期間表記 → リンクの順に並ぶ）
_SPOT_PAGE = """
<a href="/m/tvdivq00000014fy-att/Sep2026_j.pdf">月間</a>
<td>2026年9月第4週(9月24日～9月25日)</td>
<a href="/m/t13vrt0000026acb-att/stock_1_w_20260924_20260925.pdf">PDF</a>
<a href="/m/t13vrt0000026acb-att/stock_1_w_20260924_20260925.xlsx">Excel</a>
<td>2026年9月第3週(9月14日～9月18日)</td>
<a href="/m/t13vrt000002644i-att/stock_1_w_20260914_20260918.xlsx">Excel</a>
<td>2026年9月第2週(9月7日～9月11日)</td>
<a href="/m/t13vrt000001yhs9-att/stock_vol_1_260902.xls">株数</a>
<a href="/m/t13vrt000001yhs9-att/stock_val_1_260902.xls">金額</a>
<td>2026年9月第1週(8月31日～9月4日)</td>
<a href="/m/t13vrt000001y0ho-att/stock_val_1_260901.xls">金額</a>
<a href="/m/tvdivq00000014fy-att/revision_information_j.xls">訂正</a>
<a href="/m/tvdivq00000014fy-att/stock_1_w_YYYYMMDD_YYYYMMDD.xlsx">見本</a>
"""


def test_spot_files_by_week_reads_both_formats():
    files = spot_files_by_week(_SPOT_PAGE)
    assert sorted(files) == [date(2026, 9, 4), date(2026, 9, 11), date(2026, 9, 18), date(2026, 9, 25)]
    assert files[date(2026, 9, 25)].endswith("stock_1_w_20260924_20260925.xlsx")
    assert files[date(2026, 9, 11)].endswith("stock_val_1_260902.xls")   # 株数(stock_vol)ではない
    assert max(files) == date(2026, 9, 25)


def test_unreported_weeks_detects_skipped_week():
    """10/2 の実行: 9/18週と9/25週が両方掲載済み → 最新(9/25)より古い 9/18週 の欠落を検出する。"""
    published = [date(2026, 9, 4), date(2026, 9, 11), date(2026, 9, 18), date(2026, 9, 25)]
    reported = {date(2026, 9, 4), date(2026, 9, 11)}
    assert unreported_weeks(published, date(2026, 9, 25), reported.__contains__) == [date(2026, 9, 18)]
    reported.add(date(2026, 9, 18))
    assert unreported_weeks(published, date(2026, 9, 25), reported.__contains__) == []
